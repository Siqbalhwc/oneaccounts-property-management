"""
Excel upload of monthly bill-basis charges (electricity, water, gas, ...).

FLOW (all of it runs here in the backend; the browser only shows results)
  1. build_template()  -- an .xlsx pre-filled with every active lease's
                          building, room, tenant and its own utility names;
                          the user only types amounts.
  2. evaluate_upload() -- reads the uploaded file and decides, row by row,
                          exactly what WOULD happen. Writes nothing.
  3. commit_upload()   -- repeats step 2 against fresh data, then saves the
                          valid rows and recalculates the affected invoices.

SAFETY
  * Every query goes through the caller's own RLS-scoped Supabase client, so
    a row can only ever match a lease in the caller's company (and, for a
    building-scoped user, only their buildings). Rooms outside that simply
    "aren't found". The company is never read from the file.
  * Only a draft or sent invoice is ever recalculated. An invoice with a
    payment recorded is reported as "locked" and left completely alone.
  * Nothing here edits a posted journal entry: the invoice recalculation
    reuses invoicing.resync_invoice_for_month, which reverses and reposts.
  * A row that cannot be matched leaves its lease exactly as it was.
  * Saving is idempotent: the same file can be uploaded again safely; a bill
    for the same lease + month + utility is replaced, not duplicated.
"""

import math
import re
import time
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO
from typing import Any, Optional

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from supabase import Client

from app.crud.generic import write_audit_log
from app.services.invoice_batch import fetch_all, fetch_in
from app.services.invoicing import EDITABLE_INVOICE_STATUSES, merge_current_charges, resync_invoice_for_month
from app.services.utility_bills import fetch_bills, is_utility_label, label_key

SHEET_NAME = "Utility bills"
MAX_ROWS = 5000
MAX_BYTES = 5 * 1024 * 1024
COMMIT_TIME_BUDGET = 18.0  # seconds; Vercel kills the function at 30

READ_COLUMNS = [
    "building_name",
    "room_number",
    "utility",
    "bill_month",
    "amount",
    "bill_reference",
    "previous_reading",
    "current_reading",
]
INFO_COLUMNS = ["tenant", "current_lease_amount"]
TEMPLATE_COLUMNS = ["building_name", "room_number", "tenant", "utility", "bill_month", "amount", "bill_reference", "previous_reading", "current_reading", "current_lease_amount"]
REQUIRED_COLUMNS = ["building_name", "room_number", "utility", "amount"]
HEADER_ALIASES = {
    "building": "building_name",
    "room": "room_number",
    "room_no": "room_number",
    "label": "utility",
    "charge": "utility",
    "bill": "amount",
    "bill_amount": "amount",
    "month": "bill_month",
    "reference": "bill_reference",
    "ref": "bill_reference",
}

MONTH_NAMES = {
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
}


# --------------------------------------------------------------------------
# Small parsing helpers
# --------------------------------------------------------------------------
def norm_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value).split())


def norm_key(value: Any) -> str:
    return norm_text(value).lower()


def month_start(d: date) -> date:
    return date(d.year, d.month, 1)


def parse_month_value(value: Any) -> Optional[date]:
    """Blank -> None. Accepts a real Excel date, 2026-10, 2026-10-15, 10/2026,
    'October 2026' or 'Oct 2026'. Anything else raises ValueError."""
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (datetime, date)):
        return date(value.year, value.month, 1)
    s = str(value).strip()
    m = re.match(r"^(\d{4})-(\d{1,2})(?:-\d{1,2})?$", s)
    if m and 1 <= int(m.group(2)) <= 12:
        return date(int(m.group(1)), int(m.group(2)), 1)
    m = re.match(r"^(\d{1,2})/(\d{4})$", s)
    if m and 1 <= int(m.group(1)) <= 12:
        return date(int(m.group(2)), int(m.group(1)), 1)
    m = re.match(r"^([A-Za-z]+)[ ,.-]*(\d{4})$", s)
    if m:
        name = m.group(1).lower()
        for full, num in MONTH_NAMES.items():
            if name == full or (len(name) >= 3 and full.startswith(name)):
                return date(int(m.group(2)), num, 1)
    raise ValueError(f'Could not read the month "{s}". Use a format like 2026-10.')


def parse_number(value: Any, what: str) -> Optional[float]:
    """Blank -> None. Otherwise a non-negative number rounded to 2 places."""
    if value is None or (isinstance(value, str) and value.strip() == ""):
        return None
    if isinstance(value, bool):
        raise ValueError(f"{what} must be a number.")
    if isinstance(value, (int, float)):
        n = float(value)
    else:
        text = re.sub(r"(?i)\brs\.?\b|pkr", "", str(value)).replace(",", "").strip()
        try:
            n = float(text)
        except ValueError:
            raise ValueError(f"{what} must be a number.")
    if not math.isfinite(n):
        raise ValueError(f"{what} must be a number.")
    if n < 0:
        raise ValueError(f"{what} cannot be negative.")
    if n > 1_000_000_000:
        raise ValueError(f"{what} is unrealistically large.")
    return round(n, 2)


def _natural(value: str):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", value or "")]


# --------------------------------------------------------------------------
# Template
# --------------------------------------------------------------------------
def build_template(supabase: Client, bill_month: date, building_id: Optional[str]) -> bytes:
    buildings = fetch_all(lambda: supabase.table("buildings").select("id, name").order("id"))
    building_name = {b["id"]: b["name"] for b in buildings}
    rooms_q = lambda: supabase.table("rooms").select("id, room_number, building_id").order("id")
    if building_id:
        rooms = fetch_all(lambda: rooms_q().eq("building_id", building_id))
    else:
        rooms = fetch_all(rooms_q)
    room_by_id = {r["id"]: r for r in rooms}

    leases = fetch_in(
        supabase, "leases", "id, room_id, tenant_id, status", "room_id", list(room_by_id),
        extra=lambda q: q.eq("status", "active"),
    )
    lease_ids = [l["id"] for l in leases]
    tenants = {t["id"]: t["full_name"] for t in fetch_in(supabase, "tenants", "id, full_name", "id", list({l["tenant_id"] for l in leases}))}
    charges_by_lease: dict = defaultdict(list)
    for c in fetch_in(supabase, "lease_charges", "*", "lease_id", lease_ids):
        charges_by_lease[c["lease_id"]].append(c)
    existing = fetch_bills(supabase, lease_ids, bill_month)

    rows = []
    for lease in leases:
        room = room_by_id.get(lease["room_id"])
        if not room:
            continue
        bills = {b["label_key"]: float(b["amount"]) for b in existing.get(lease["id"], [])}
        for charge in merge_current_charges(charges_by_lease.get(lease["id"], [])):
            if not is_utility_label(charge["label"]):
                continue
            rows.append(
                (
                    building_name.get(room["building_id"], ""),
                    str(room["room_number"]),
                    tenants.get(lease["tenant_id"], ""),
                    charge["label"],
                    float(charge["amount"]),
                    bills.get(label_key(charge["label"])),
                )
            )
    rows.sort(key=lambda r: (_natural(r[0]), _natural(r[1]), r[3].lower()))

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(TEMPLATE_COLUMNS)
    head_fill = PatternFill("solid", fgColor="0F5C63")
    info_fill = PatternFill("solid", fgColor="7A8790")
    for idx, cell in enumerate(ws[1], start=1):
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = info_fill if TEMPLATE_COLUMNS[idx - 1] in INFO_COLUMNS else head_fill
        cell.alignment = Alignment(vertical="center")
    month_text = bill_month.strftime("%Y-%m")
    for b, r, tenant, utility, lease_amt, bill in rows:
        ws.append([b, r, tenant, utility, month_text, bill, None, None, None, lease_amt])
    for col, width in zip("ABCDEFGHIJ", (22, 14, 26, 22, 12, 14, 18, 16, 16, 20)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = "@"

    guide = wb.create_sheet("Instructions")
    for line in [
        "How to use this sheet",
        "1. Type the bill in the 'amount' column for each row. Leave 'amount' blank to leave that row unchanged.",
        "2. Do not change building_name, room_number, utility or bill_month. The grey columns are for your information only.",
        "3. To bill a utility the lease does not charge yet, add a row with the building, room, utility name and amount.",
        "4. An amount of 0 removes that charge from the month's invoice.",
        "5. bill_reference, previous_reading and current_reading are optional.",
        "6. Uploading the same month again replaces the earlier amounts.",
        "7. Invoices that already have a payment recorded are never changed.",
    ]:
        guide.append([line])
    guide.column_dimensions["A"].width = 110
    guide["A1"].font = Font(bold=True)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Reading the uploaded file
# --------------------------------------------------------------------------
def read_upload(content: bytes) -> list[dict]:
    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    if len(content) > MAX_BYTES:
        raise HTTPException(status_code=400, detail="That file is too large (limit 5 MB). Split it by building.")
    try:
        wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Could not read that file. Please upload the .xlsx template.")
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb[wb.sheetnames[0]]
    raw = list(ws.iter_rows(values_only=True))
    if not raw:
        raise HTTPException(status_code=400, detail="The file has no rows.")

    headers = []
    for h in raw[0]:
        key = norm_key(h).replace(" ", "_").replace("-", "_")
        headers.append(HEADER_ALIASES.get(key, key))
    missing = [c for c in REQUIRED_COLUMNS if c not in headers]
    if missing:
        raise HTTPException(
            status_code=400,
            detail="The file is missing these columns: " + ", ".join(missing) + ". Please use the downloaded template.",
        )
    out = []
    for offset, values in enumerate(raw[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        out.append({"_row": offset, **{h: v for h, v in zip(headers, values) if h}})
        if len(out) > MAX_ROWS:
            raise HTTPException(status_code=400, detail=f"Too many rows (limit {MAX_ROWS}). Split the file by building.")
    return out


# --------------------------------------------------------------------------
# Evaluation (no writes)
# --------------------------------------------------------------------------
def evaluate_upload(supabase: Client, rows: list[dict], bill_month: date) -> dict:
    """
    Decides what would happen to each row. Returns
      {"rows": [...], "groups": {lease_id: {...}}, "summary": {...}}
    Row outcomes: update | stored | locked | error | blank.
    """
    wanted_buildings = {norm_key(r.get("building_name")) for r in rows if norm_text(r.get("building_name"))}
    buildings = fetch_all(lambda: supabase.table("buildings").select("id, name").order("id"))
    buildings_by_key: dict = defaultdict(list)
    for b in buildings:
        if norm_key(b["name"]) in wanted_buildings:
            buildings_by_key[norm_key(b["name"])].append(b)
    building_ids = [b["id"] for lst in buildings_by_key.values() for b in lst]

    rooms = fetch_in(supabase, "rooms", "id, room_number, building_id", "building_id", building_ids)
    rooms_by_key: dict = defaultdict(list)
    for r in rooms:
        rooms_by_key[(r["building_id"], norm_key(r["room_number"]))].append(r)

    leases = fetch_in(
        supabase, "leases", "*", "room_id", [r["id"] for r in rooms], extra=lambda q: q.eq("status", "active")
    )
    leases_by_room: dict = defaultdict(list)
    for l in leases:
        leases_by_room[l["room_id"]].append(l)
    lease_ids = [l["id"] for l in leases]

    tenants = {t["id"]: t["full_name"] for t in fetch_in(supabase, "tenants", "id, full_name", "id", list({l["tenant_id"] for l in leases}))}
    charges_by_lease: dict = defaultdict(list)
    for c in fetch_in(supabase, "lease_charges", "*", "lease_id", lease_ids):
        charges_by_lease[c["lease_id"]].append(c)
    invoices = {
        i["lease_id"]: i
        for i in fetch_in(
            supabase, "invoices", "id, lease_id, status, total_amount, invoice_month", "lease_id", lease_ids,
            extra=lambda q: q.eq("invoice_month", str(bill_month)),
        )
    }
    existing_bills = fetch_bills(supabase, lease_ids, bill_month)
    mapped_labels = {
        m["label"] for m in fetch_all(lambda: supabase.table("charge_type_accounts").select("id, label").order("id"))
    }
    building_name_by_id = {b["id"]: b["name"] for b in buildings}
    room_by_id = {r["id"]: r for r in rooms}

    out_rows: list[dict] = []
    groups: dict = {}
    seen: dict = {}
    for src in rows:
        n = src["_row"]
        entry: dict = {
            "row": n,
            "building": norm_text(src.get("building_name")),
            "room": norm_text(src.get("room_number")),
            "tenant": "",
            "utility": norm_text(src.get("utility")),
            "lease_amount": None,
            "amount": None,
            "invoice_status": None,
            "outcome": "error",
            "message": "",
            "warnings": [],
        }
        out_rows.append(entry)

        def fail(msg: str) -> None:
            entry["outcome"] = "error"
            entry["message"] = msg

        try:
            amount = parse_number(src.get("amount"), "Amount")
            prev_reading = parse_number(src.get("previous_reading"), "Previous reading")
            curr_reading = parse_number(src.get("current_reading"), "Current reading")
            row_month = parse_month_value(src.get("bill_month"))
        except ValueError as e:
            fail(str(e))
            continue
        if amount is None:
            entry["outcome"] = "blank"
            entry["message"] = "No amount entered, left unchanged."
            continue
        entry["amount"] = amount
        if not entry["building"] or not entry["room"] or not entry["utility"]:
            fail("Building, room and utility are all required.")
            continue
        if row_month is not None and row_month != bill_month:
            fail(f"This row is for {row_month.strftime('%B %Y')}, but you selected {bill_month.strftime('%B %Y')}.")
            continue

        blds = buildings_by_key.get(norm_key(entry["building"]), [])
        if not blds:
            fail("Building not found. The lease stays as it is.")
            continue
        candidate_rooms = [r for b in blds for r in rooms_by_key.get((b["id"], norm_key(entry["room"])), [])]
        if not candidate_rooms:
            fail("Room not found in that building. The lease stays as it is.")
            continue
        room_leases = [l for r in candidate_rooms for l in leases_by_room.get(r["id"], [])]
        if not room_leases:
            fail("No active lease on this room. Nothing changes.")
            continue
        if len(room_leases) > 1:
            fail("More than one active lease on this room. Fix the room's leases first.")
            continue
        lease = room_leases[0]
        entry["tenant"] = tenants.get(lease["tenant_id"], "")

        key = label_key(entry["utility"])
        current = merge_current_charges(charges_by_lease.get(lease["id"], []))
        matches = [c for c in current if label_key(c["label"]) == key]
        match_labels = sorted({c["label"] for c in matches})
        if len(match_labels) > 1:
            fail("This lease has more than one charge for this utility (" + ", ".join(match_labels) + "). Please tidy the lease first.")
            continue
        if matches:
            resolved_label = matches[0]["label"]
            entry["lease_amount"] = float(matches[0]["amount"])
        else:
            resolved_label = entry["utility"]
            if resolved_label not in mapped_labels:
                entry["warnings"].append(
                    f'No income account is set up for "{resolved_label}", so it will post to Other Income (4100).'
                )
        entry["utility"] = resolved_label

        dup_key = (lease["id"], key)
        if dup_key in seen:
            fail(f"Duplicate of row {seen[dup_key]} (same room and utility).")
            continue
        seen[dup_key] = n

        previous = next((b for b in existing_bills.get(lease["id"], []) if b["label_key"] == key), None)
        if previous is not None:
            entry["warnings"].append(f"Replaces the bill saved earlier ({float(previous['amount']):,.2f}).")
        if amount == 0:
            entry["warnings"].append("0 removes this charge from the month's invoice.")

        invoice = invoices.get(lease["id"])
        entry["invoice_status"] = invoice["status"] if invoice else None
        if invoice is None:
            entry["outcome"] = "stored"
            entry["message"] = "Saved. It will be used when this month's invoice is generated."
        elif invoice["status"] in EDITABLE_INVOICE_STATUSES:
            entry["outcome"] = "update"
            entry["message"] = "The invoice will be recalculated."
        else:
            entry["outcome"] = "locked"
            entry["message"] = "Not applied. A payment is already recorded on this invoice."
            continue

        room = room_by_id.get(lease["room_id"], {})
        group = groups.setdefault(
            lease["id"],
            {
                "lease": lease,
                "invoice": invoice,
                "label": f"{building_name_by_id.get(room.get('building_id'), '')} / {room.get('room_number', '')}",
                "first_row": n,
                "items": [],
            },
        )
        group["items"].append(
            {
                "row": n,
                "label": resolved_label,
                "label_key": key,
                "amount": amount,
                "bill_reference": norm_text(src.get("bill_reference")) or None,
                "previous_reading": prev_reading,
                "current_reading": curr_reading,
            }
        )

    summary = {"rows_checked": len(out_rows), "update": 0, "stored": 0, "locked": 0, "error": 0, "blank": 0}
    for r in out_rows:
        summary[r["outcome"]] += 1
    return {"rows": out_rows, "groups": groups, "summary": summary}


def public_rows(rows: list[dict]) -> list[dict]:
    return [dict(r) for r in rows]


# --------------------------------------------------------------------------
# Commit
# --------------------------------------------------------------------------
def commit_upload(
    supabase: Client,
    company_id: str,
    user_id: str,
    rows: list[dict],
    bill_month: date,
    start: int = 0,
    time_budget: float = COMMIT_TIME_BUDGET,
) -> dict:
    """
    Saves the valid rows and recalculates the affected draft/sent invoices.
    Time-budgeted: leases are handled in a stable order; anything not reached
    comes back as next_start so the caller simply asks again (idempotent).
    """
    started = time.monotonic()
    ev = evaluate_upload(supabase, rows, bill_month)
    order = sorted(ev["groups"].items(), key=lambda kv: kv[1]["first_row"])
    outcomes: dict[int, dict] = {}
    next_start: Optional[int] = None

    for pos in range(start, len(order)):
        if time.monotonic() - started > time_budget:
            next_start = pos
            break
        lease_id, group = order[pos]
        _commit_group(supabase, company_id, user_id, lease_id, group, bill_month, outcomes)

    return {
        "rows": ev["rows"],
        "outcomes": [{"row": r, **o} for r, o in sorted(outcomes.items())],
        "summary": ev["summary"],
        "total_groups": len(order),
        "next_start": next_start,
    }


def _commit_group(
    supabase: Client,
    company_id: str,
    user_id: str,
    lease_id: str,
    group: dict,
    bill_month: date,
    outcomes: dict,
) -> None:
    items = group["items"]

    def mark(outcome: str, message: str, total: Optional[float] = None) -> None:
        for it in items:
            outcomes[it["row"]] = {"outcome": outcome, "message": message, "invoice_total": total}

    try:
        # Re-check the invoice right before writing: it may have been paid
        # (or generated) since the file was checked.
        current = (
            supabase.table("invoices")
            .select("id, status, total_amount")
            .eq("lease_id", lease_id)
            .eq("invoice_month", str(bill_month))
            .execute()
            .data
        )
        invoice = current[0] if current else None
        if invoice is not None and invoice["status"] not in EDITABLE_INVOICE_STATUSES:
            mark("locked", "Not applied. A payment is already recorded on this invoice.")
            return

        payload = [
            {
                "company_id": company_id,
                "lease_id": lease_id,
                "bill_month": str(bill_month),
                "label": it["label"],
                "label_key": it["label_key"],
                "amount": it["amount"],
                "bill_reference": it["bill_reference"],
                "previous_reading": it["previous_reading"],
                "current_reading": it["current_reading"],
                "source": "excel_upload",
                "uploaded_by": user_id,
            }
            for it in items
        ]
        saved = supabase.table("utility_bills").upsert(payload, on_conflict="lease_id,bill_month,label_key").execute().data
        for s in saved or []:
            write_audit_log(supabase, company_id, user_id, "update", "utility_bills", s["id"], after={"label": s["label"], "bill_month": s["bill_month"], "amount": s["amount"]})

        if invoice is None:
            mark("stored", "Saved. It will be used when this month's invoice is generated.")
            return

        old_total = float(invoice["total_amount"])
        updated = resync_invoice_for_month(supabase, company_id, group["lease"], bill_month)
        if updated is None:
            mark("failed", "The bill was saved but the invoice could not be recalculated. Upload again to retry.")
            return
        new_total = float(updated["total_amount"])
        write_audit_log(
            supabase, company_id, user_id, "update", "invoices", invoice["id"],
            before={"total_amount": old_total}, after={"total_amount": new_total},
        )
        mark("updated", f"Invoice recalculated: {old_total:,.2f} to {new_total:,.2f}.", new_total)
    except Exception as e:
        message = getattr(e, "detail", None) or getattr(e, "message", None) or str(e)
        mark("failed", f"Could not finish this lease ({str(message)[:200]}). Upload the file again to retry.")
